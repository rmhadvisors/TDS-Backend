import io
import json
import os
from decimal import Decimal
import pytest
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def test_health_endpoint():
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "ok"


def test_compare_endpoint_flow():
    # Sample original data
    master_old = json.dumps([
        {"GUID": "M-1", "NAME": "Aayat Enterprises", "PARENT": "Sundry Creditors", "PAN": "ABCDE1234F", "ISTDSAPPLICABLE": "Yes", "TDSSECTION": "194C"}
    ]).encode("utf-8")

    tx_old = json.dumps([
        {
            "GUID": "TX-101",
            "VOUCHERNUMBER": "101",
            "VOUCHERTYPENAME": "Payment",
            "DATE": "2026-08-01",
            "PARTYLEDGERNAME": "Aayat Enterprises",
            "ALLLEDGERENTRIES.LIST": [
                {"LEDGERNAME": "Aayat Enterprises", "AMOUNT": "-90000.00", "ISDEEMEDPOSITIVE": "No"},
                {"LEDGERNAME": "TDS on Contractors", "AMOUNT": "-10000.00", "ISDEEMEDPOSITIVE": "No", "IS_TDS": "Yes"},
                {"LEDGERNAME": "Contractor Charges", "AMOUNT": "100000.00", "ISDEEMEDPOSITIVE": "Yes"}
            ]
        }
    ]).encode("utf-8")

    master_new = master_old
    tx_new = tx_old

    files = {
        "master_old": ("Master.json", io.BytesIO(master_old), "application/json"),
        "transactions_old": ("Transactions.json", io.BytesIO(tx_old), "application/json"),
        "master_new": ("Master1.json", io.BytesIO(master_new), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(tx_new), "application/json"),
    }

    res = client.post("/api/compare", files=files)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["success"] is True
    assert data["summary"]["original_transactions"] == 1
    assert data["summary"]["updated_transactions"] == 1
    assert data["summary"]["unchanged"] == 1
    assert data["reconciliation"]["is_valid"] is True
    assert data["reconciliation"]["transaction_count"] == "PASS"

    download_id = data["report"]["download_id"]
    assert data["report"]["party_name"] == "Aayat Enterprises"
    assert data["report"]["period"] == "01/08/2026"

    dl_res = client.get(f"/api/download/{download_id}")
    assert dl_res.status_code == 200
    assert dl_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(dl_res.content) > 1000

    # Verify Party Name and Period in Excel file
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(dl_res.content))
    ws = wb["Summary"]
    assert ws["A3"].value == "Party Name"
    assert ws["B3"].value == "Aayat Enterprises"
    assert ws["A4"].value == "Period"
    assert ws["B4"].value == "01/08/2026"


def _single_master_payload(party_ledger):
    master = json.dumps([
        {"GUID": "M-1", "NAME": party_ledger, "PARENT": "Sundry Creditors", "PAN": "ABCDE1234F", "ISTDSAPPLICABLE": "Yes", "TDSSECTION": "194C"},
        {"GUID": "M-2", "NAME": "TDS on Contractors", "PARENT": "Duties & Taxes"},
        {"GUID": "M-3", "NAME": "Contractor Charges", "PARENT": "Indirect Expenses"},
    ]).encode("utf-8")
    tx = json.dumps([{
        "GUID": "TX-101", "VOUCHERNUMBER": "101", "VOUCHERTYPENAME": "Payment", "DATE": "2026-08-01",
        "PARTYLEDGERNAME": "Aayat Enterprises",
        "ALLLEDGERENTRIES.LIST": [
            {"LEDGERNAME": "Aayat Enterprises", "AMOUNT": "-90000.00", "ISDEEMEDPOSITIVE": "No"},
            {"LEDGERNAME": "TDS on Contractors", "AMOUNT": "-10000.00", "ISDEEMEDPOSITIVE": "No", "IS_TDS": "Yes"},
            {"LEDGERNAME": "Contractor Charges", "AMOUNT": "100000.00", "ISDEEMEDPOSITIVE": "Yes"},
        ],
    }]).encode("utf-8")
    return {
        "transactions_old": ("Transactions.json", io.BytesIO(tx), "application/json"),
        "master_new": ("Master.json", io.BytesIO(master), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(tx), "application/json"),
    }


def test_compare_with_latest_master_only():
    # No original master: the latest one is used for both snapshots
    res = client.post("/api/compare", files=_single_master_payload("Aayat Enterprises"))
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["summary"]["original_transactions"] == 1
    assert data["summary"]["unchanged"] == 1
    assert data["reconciliation"]["is_valid"] is True
    assert data["missing_master_ledgers"] == []


def test_latest_master_missing_a_filed_party_is_reported():
    # The party of the filed voucher was renamed in the latest master
    res = client.post("/api/compare", files=_single_master_payload("Aayat Enterprises Pvt Ltd"))
    assert res.status_code == 200, res.text
    assert res.json()["missing_master_ledgers"] == ["Aayat Enterprises"]


def test_typed_period_is_shown_in_summary():
    res = client.post("/api/compare", files=_single_master_payload("Aayat Enterprises"),
                      data={"period_from": "2026-04-01", "period_to": "2027-03-31"})
    assert res.status_code == 200, res.text
    assert res.json()["report"]["period"] == "01/04/2026 to 31/03/2027"

    import openpyxl
    content = client.get(f"/api/download/{res.json()['report']['download_id']}").content
    assert openpyxl.load_workbook(io.BytesIO(content))["Summary"]["B4"].value == "01/04/2026 to 31/03/2027"


def test_period_end_not_typed_comes_from_vouchers():
    # Only From typed: To is the latest voucher date (01/08/2026)
    res = client.post("/api/compare", files=_single_master_payload("Aayat Enterprises"), data={"period_from": "2026-04-01"})
    assert res.status_code == 200, res.text
    assert res.json()["report"]["period"] == "01/04/2026 to 01/08/2026"


def test_compare_with_custom_party_and_period():
    master_json = json.dumps([{"GUID": "M-1", "NAME": "Custom Vendor Ltd", "ISTDSAPPLICABLE": "Yes"}]).encode("utf-8")
    tx_json = json.dumps([{
        "GUID": "TX-1",
        "VOUCHERNUMBER": "1",
        "VOUCHERTYPENAME": "Payment",
        "DATE": "2025-04-15",
        "PARTYLEDGERNAME": "Custom Vendor Ltd",
        "ALLLEDGERENTRIES.LIST": [
            {"LEDGERNAME": "Custom Vendor Ltd", "AMOUNT": "-50000.00", "ISDEEMEDPOSITIVE": "No"},
            {"LEDGERNAME": "Expense", "AMOUNT": "50000.00", "ISDEEMEDPOSITIVE": "Yes"}
        ]
    }]).encode("utf-8")

    files = {
        "master_old": ("Master.json", io.BytesIO(master_json), "application/json"),
        "transactions_old": ("Transactions.json", io.BytesIO(tx_json), "application/json"),
        "master_new": ("Master1.json", io.BytesIO(master_json), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(tx_json), "application/json"),
    }
    data_payload = {
        "party_name": "ABC Corporation Ltd",
        "period": "FY 2025-26 (Q1)",
    }

    res = client.post("/api/compare", files=files, data=data_payload)
    assert res.status_code == 200, res.text
    resp_data = res.json()
    assert resp_data["report"]["party_name"] == "ABC Corporation Ltd"
    assert resp_data["report"]["period"] == "FY 2025-26 (Q1)"

    download_id = resp_data["report"]["download_id"]
    dl_res = client.get(f"/api/download/{download_id}")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(dl_res.content))
    ws = wb["Summary"]
    assert ws["A3"].value == "Party Name"
    assert ws["B3"].value == "ABC Corporation Ltd"
    assert ws["A4"].value == "Period"
    assert ws["B4"].value == "FY 2025-26 (Q1)"


def test_compare_with_sample_files():
    sample_dir = os.path.join(os.path.dirname(__file__), "..", "..", "sample_data")
    with open(os.path.join(sample_dir, "Master.json"), "rb") as f:
        m_old = f.read()
    with open(os.path.join(sample_dir, "Transactions.json"), "rb") as f:
        t_old = f.read()
    with open(os.path.join(sample_dir, "Master1.json"), "rb") as f:
        m_new = f.read()
    with open(os.path.join(sample_dir, "Transactions1.json"), "rb") as f:
        t_new = f.read()

    files = {
        "master_old": ("Master.json", io.BytesIO(m_old), "application/json"),
        "transactions_old": ("Transactions.json", io.BytesIO(t_old), "application/json"),
        "master_new": ("Master1.json", io.BytesIO(m_new), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(t_new), "application/json"),
    }

    res = client.post("/api/compare", files=files)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["success"] is True
    summary = data["summary"]
    assert summary["original_transactions"] == 7
    assert summary["updated_transactions"] == 6
    assert summary["unchanged"] == 6
    assert summary["missing"] == 1
    assert summary["modified"] == 0
    assert summary["recreated"] == 0
    assert summary["new"] == 0
    assert summary["unresolved"] == 0

    assert data["reconciliation"]["is_valid"] is True
    assert data["reconciliation"]["transaction_count"] == "PASS"
    assert data["reconciliation"]["amount"] == "PASS"
    assert data["reconciliation"]["tds"] == "PASS"


def test_compare_with_utf16_and_bom_encodings():
    master_json = json.dumps([
        {"GUID": "M-1", "NAME": "Aayat Enterprises", "PARENT": "Sundry Creditors", "PAN": "ABCDE1234F", "ISTDSAPPLICABLE": "Yes", "TDSSECTION": "194C"}
    ])

    tx_json = json.dumps([
        {
            "GUID": "TX-101",
            "VOUCHERNUMBER": "101",
            "VOUCHERTYPENAME": "Payment",
            "DATE": "2026-08-01",
            "PARTYLEDGERNAME": "Aayat Enterprises",
            "ALLLEDGERENTRIES.LIST": [
                {"LEDGERNAME": "Aayat Enterprises", "AMOUNT": "-90000.00", "ISDEEMEDPOSITIVE": "No"},
                {"LEDGERNAME": "TDS on Contractors", "AMOUNT": "-10000.00", "ISDEEMEDPOSITIVE": "No", "IS_TDS": "Yes"},
                {"LEDGERNAME": "Contractor Charges", "AMOUNT": "100000.00", "ISDEEMEDPOSITIVE": "Yes"}
            ]
        }
    ])

    # Test UTF-16 LE (with BOM), UTF-16 BE, and UTF-8 with BOM
    files = {
        "master_old": ("Master.json", io.BytesIO(master_json.encode("utf-16")), "application/json"),
        "transactions_old": ("Transactions.json", io.BytesIO(tx_json.encode("utf-16-le")), "application/json"),
        "master_new": ("Master1.json", io.BytesIO(master_json.encode("utf-8-sig")), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(tx_json.encode("utf-16-be")), "application/json"),
    }

    res = client.post("/api/compare", files=files)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["success"] is True
    assert data["summary"]["original_transactions"] == 1


def test_corrupted_json_error():
    files = {
        "master_old": ("Master.json", io.BytesIO(b"INVALID JSON {{{{"), "application/json"),
        "transactions_old": ("Transactions.json", io.BytesIO(b"[]"), "application/json"),
        "master_new": ("Master1.json", io.BytesIO(b"[]"), "application/json"),
        "transactions_new": ("Transactions1.json", io.BytesIO(b"[]"), "application/json"),
    }
    res = client.post("/api/compare", files=files)
    assert res.status_code == 400
    assert "error" in res.json()
